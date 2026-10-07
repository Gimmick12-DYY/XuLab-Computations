#!/usr/bin/env python3
"""Fill TF1000cells_known_motifs (Full) yellow Codebook-only rows from AME results."""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path("/work/users/d/y/dyy12/XuLab")
MA = ROOT / "motif_analysis"
XLSX = MA / "TF1000cells_motifs.xlsx"
PEAKS = MA / "motif"
BINS = MA / "motif_bins"

RESCUED = [
    "ZNF507", "MBD3", "ZNF703", "GLYR1", "FLYWCH1", "ZNF367",
    "ZNF746", "POGK", "ZBTB40", "ZNF606", "ZNF503",
]


def tok(name: str) -> str:
    return re.split(r"[(/_.]", name.strip())[0].lower()


def ame_own(tf: str, root: Path, bg: str) -> str:
    """adj_p of own Codebook motif, or F if ran but absent, or '' if not run."""
    import csv

    base = root / tf.lower()
    tagged = base / f"ame_{bg}" / tf.lower() / "ame.tsv"
    legacy = base / "ame" / tf.lower() / "ame.tsv"
    f = tagged if tagged.is_file() else (legacy if bg == "genome" and legacy.is_file() else None)
    if f is None:
        return ""
    want = {tf.lower()}
    best = None
    with open(f) as fh:
        for r in csv.DictReader(fh, delimiter="\t"):
            if not (r.get("rank", "").strip().isdigit()):
                continue
            mid = r.get("motif_ID", "") or r.get("motif_alt_ID", "")
            if tok(mid) not in want:
                continue
            raw = (r.get("adj_p-value") or "").strip()
            try:
                v = float(raw)
            except ValueError:
                continue
            if best is None or v < best[0]:
                best = (v, raw)
    if best is None:
        return "F"
    v, raw = best
    if v > 0:
        return f"{v:.2E}"
    return raw.upper()


def main() -> int:
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["TF1000cells_known_motifs (Full)"]
    # map TF -> row
    rows = {}
    for i in range(1, ws.max_row + 1):
        v = ws.cell(i, 1).value
        if v and str(v).upper() in RESCUED:
            rows[str(v).upper()] = i

    missing = [t for t in RESCUED if t not in rows]
    if missing:
        print(f"[warn] TFs not found in Full sheet: {missing}", file=sys.stderr)

    updated = 0
    for tf in RESCUED:
        i = rows.get(tf)
        if not i:
            continue
        # cols: 3 Raw AME peaks, 4 Homer peaks, 5 AME bins, 6 Homer bins
        #       7 Imp genome AME peaks, 8 Homer, 9 AME bins, 10 Homer
        #       11 Imp OCR AME peaks, 12 Homer, 13 AME bins, 14 Homer
        #       15 HOCOMOCO, 16 Homer db, 17 Codebook
        gap = ame_own(tf, PEAKS, "genome")
        oap = ame_own(tf, PEAKS, "ocr")
        gab = ame_own(tf, BINS, "genome")
        oab = ame_own(tf, BINS, "ocr")

        # Raw stays NA (no codebook raw run); Homer stays NA (not in HOMER lib)
        for c in (3, 4, 5, 6, 8, 10, 12, 14):
            ws.cell(i, c, "NA")
        ws.cell(i, 7, gap if gap else "NA")
        ws.cell(i, 9, gab if gab else "NA")
        ws.cell(i, 11, oap if oap else "NA")
        ws.cell(i, 13, oab if oab else "NA")
        ws.cell(i, 15, "#N/A")
        ws.cell(i, 16, "#N/A")
        ws.cell(i, 17, tf)
        print(f"{tf}: genome_peaks={gap or 'NA'} ocr_peaks={oap or 'NA'} "
              f"genome_bins={gab or 'NA'} ocr_bins={oab or 'NA'}")
        updated += 1

    wb.save(XLSX)
    print(f"[done] updated {updated} rows in {XLSX}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
